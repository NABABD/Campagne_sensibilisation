"""
Génère un rapport Excel lisible à partir des résultats d'une campagne GoPhish.
"""

from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from gophish_client import GoPhishClient


STATUS_LABELS = {
    "Sent": "Email envoyé",
    "Email Opened": "Email ouvert",
    "Clicked Link": "A cliqué sur le lien",
    "Submitted Data": "A saisi des informations",
    "Email Reported": "A signalé le mail",
}


def build_excel_report(campaign_id: int, output_path: str):
    client = GoPhishClient()
    results = client.get_campaign_results(campaign_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Résultats campagne"

    headers = ["Prénom", "Nom", "Email", "Poste", "Statut le plus avancé", "Dernière mise à jour"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C00000")

    for r in results.get("results", []):
        status = r.get("status", "")
        label = STATUS_LABELS.get(status, status)
        ws.append([
            r.get("first_name", ""),
            r.get("last_name", ""),
            r.get("email", ""),
            r.get("position", ""),
            label,
            r.get("modified_date", ""),
        ])

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = max_len + 3

    # Feuille de synthèse
    ws2 = wb.create_sheet("Synthèse")
    total = len(results.get("results", []))
    counts = {}
    for r in results.get("results", []):
        s = r.get("status", "Inconnu")
        counts[s] = counts.get(s, 0) + 1

    ws2.append(["Indicateur", "Nombre", "Pourcentage"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    ws2.append(["Total destinataires", total, "100%"])
    for status, count in counts.items():
        pct = f"{(count/total*100):.1f}%" if total else "0%"
        ws2.append([STATUS_LABELS.get(status, status), count, pct])

    wb.save(output_path)
    print(f"Rapport généré : {output_path} ({datetime.now().isoformat()})")
    return output_path
